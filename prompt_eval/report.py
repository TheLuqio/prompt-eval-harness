"""Aggregate raw results into a readable markdown report."""
from __future__ import annotations

import csv
import io
import json
from pathlib import Path


def _avg(values):
    values = [v for v in values if v is not None]
    return sum(values) / len(values) if values else None


def _pct(x):
    return "n/a" if x is None else f"{x:.0%}"


def _num(x, digits=1):
    return "n/a" if x is None else f"{x:.{digits}f}"


def summarize(rows, prompt_ids):
    out = {}
    for pid in prompt_ids:
        rs = [r for r in rows if r["prompt_id"] == pid]
        if not rs:
            continue
        per_check, per_tag = {}, {}
        for r in rs:
            for name, val in (r.get("checks") or {}).items():
                if val is not None:
                    per_check.setdefault(name, []).append(bool(val))
            for tag in r.get("tags", []):
                per_tag.setdefault(tag, []).append(bool(r["passed"]))
        out[pid] = {
            "n": len(rs),
            "pass_rate": _avg([1.0 if r["passed"] else 0.0 for r in rs]),
            "per_check": {k: _avg([1.0 if v else 0.0 for v in vs]) for k, vs in per_check.items()},
            "per_tag": {k: _avg([1.0 if v else 0.0 for v in vs]) for k, vs in per_tag.items()},
            "avg_words": _avg([r.get("words") for r in rs if not r.get("error")]),
            "avg_latency": _avg([r.get("latency_s") for r in rs if not r.get("error")]),
            "avg_in": _avg([r.get("input_tokens") for r in rs if not r.get("error")]),
            "avg_out": _avg([r.get("output_tokens") for r in rs if not r.get("error")]),
            "errors": sum(1 for r in rs if r.get("error")),
        }
    return out


def _valid_rating(value):
    try:
        number = float(str(value).strip().replace(",", "."))
    except (TypeError, ValueError):
        return None
    return number if 1 <= number <= 5 else None


def load_ratings(results_dir: Path):
    """Read the blinded manual rating sheet (xlsx, or the older csv) if the user has filled it in."""
    key_file = results_dir / "manual_rating_key.json"
    if not key_file.exists():
        return {}
    key = json.loads(key_file.read_text(encoding="utf-8"))
    pairs = []  # (blind_id, rating)
    xlsx = results_dir / "manual_rating_sheet.xlsx"
    csv_file = results_dir / "manual_rating_sheet.csv"
    if xlsx.exists():
        from openpyxl import load_workbook

        wb = load_workbook(xlsx, data_only=True)
        ws = wb["Rate emails"] if "Rate emails" in wb.sheetnames else wb.worksheets[0]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and len(row) >= 4:
                pairs.append((row[0], row[3]))
    elif csv_file.exists():
        text = csv_file.read_text(encoding="utf-8-sig")
        header = text.splitlines()[0] if text else ""
        delimiter = ";" if header.count(";") > header.count(",") else ","
        for row in csv.DictReader(io.StringIO(text), delimiter=delimiter):
            pairs.append((row.get("blind_id"), row.get("rating")))
    ratings = {}
    for blind_id, raw in pairs:
        value = _valid_rating(raw)
        info = key.get(blind_id)
        if value is not None and info:
            ratings.setdefault(info["prompt_id"], []).append(value)
    return ratings


def render_report(meta, rows, ratings=None):
    prompt_ids = meta["prompt_ids"]
    summary = summarize(rows, prompt_ids)
    p_in, p_out = meta.get("price_in"), meta.get("price_out")
    lines = []
    if meta.get("dry_run"):
        lines += [
            "> **DRY RUN: SYNTHETIC OUTPUT. These numbers come from a fake model and are not real results.**",
            "> Never quote them on a CV, README or in an interview.",
            "",
        ]
    failed_calls = sum(1 for r in rows if r.get("error"))
    if failed_calls:
        lines += [
            f"> **WARNING: {failed_calls} of {len(rows)} API calls failed. The pass rates below count failed calls as failures, so they are not meaningful. Fix the error (see Example failures) and run again.**",
            "",
        ]
    lines += [
        "# Prompt evaluation report",
        "",
        f"- Run at: {meta['run_at']}",
        f"- Model: `{meta['model']}`" + (f", temperature {meta['temperature']}" if meta.get("temperature") is not None else ""),
        f"- Leads: {meta['n_leads']} | repeats per lead: {meta['repeats']} | emails per prompt: {meta['n_leads'] * meta['repeats']}",
        "- Prompt versions: " + ", ".join(f"`{pid}` (sha {meta['prompt_sha'][pid]})" for pid in prompt_ids),
        "",
        "## Summary",
        "",
        "| Prompt | Pass rate | Injection leads passed | Avg words | Avg latency (s) | Avg tokens in/out | Est. cost per email | Errors |",
        "|---|---|---|---|---|---|---|---|",
    ]
    for pid in prompt_ids:
        s = summary.get(pid)
        if not s:
            continue
        cost = "n/a"
        if p_in is not None and p_out is not None and s["avg_in"] is not None:
            cost = f"${(s['avg_in'] * p_in + s['avg_out'] * p_out) / 1_000_000:.5f}"
        lines.append(
            f"| `{pid}` | {_pct(s['pass_rate'])} | {_pct(s['per_tag'].get('injection'))} | "
            f"{_num(s['avg_words'], 0)} | {_num(s['avg_latency'], 2)} | "
            f"{_num(s['avg_in'], 0)} / {_num(s['avg_out'], 0)} | {cost} | {s['errors']} |"
        )
    lines += [
        "",
        "Pass rate = share of emails that pass every applicable automatic check.",
        "",
        "## Pass rate by check",
        "",
    ]
    check_names = sorted({c for s in summary.values() for c in s["per_check"]})
    lines += ["| Check | " + " | ".join(f"`{p}`" for p in prompt_ids) + " |", "|---|" + "---|" * len(prompt_ids)]
    for name in check_names:
        cells = [_pct(summary[p]["per_check"].get(name)) if p in summary else "n/a" for p in prompt_ids]
        lines.append(f"| {name} | " + " | ".join(cells) + " |")
    tags = sorted({t for s in summary.values() for t in s["per_tag"]})
    lines += ["", "## Pass rate by lead type", "", "| Lead type | " + " | ".join(f"`{p}`" for p in prompt_ids) + " |", "|---|" + "---|" * len(prompt_ids)]
    for tag in tags:
        cells = [_pct(summary[p]["per_tag"].get(tag)) if p in summary else "n/a" for p in prompt_ids]
        lines.append(f"| {tag} | " + " | ".join(cells) + " |")
    if ratings:
        lines += ["", "## Manual ratings (blind, 1 to 5)", "", "| Prompt | Rated emails | Average rating |", "|---|---|---|"]
        for pid in prompt_ids:
            vals = ratings.get(pid, [])
            lines.append(f"| `{pid}` | {len(vals)} | {_num(_avg(vals), 2)} |")
    lines += ["", "## Example failures", ""]
    for pid in prompt_ids:
        failed = [r for r in rows if r["prompt_id"] == pid and not r["passed"]][:4]
        lines.append(f"**`{pid}`**" + ("" if failed else ": none"))
        for r in failed:
            if r.get("error"):
                lines.append(f"- {r['lead_id']}: API error: {r['error']}")
            else:
                bad = [k for k, v in r["checks"].items() if v is False]
                lines.append(f"- {r['lead_id']} ({', '.join(r['tags'])}): failed {', '.join(bad)}")
        lines.append("")
    lines += [
        "Full outputs for every email are in `raw_results.jsonl`.",
        "",
        "_Limits: automatic checks are heuristics; the lead set is small and synthetic; one model; "
        "manual ratings are subjective. Treat differences of a few points as noise._",
    ]
    return "\n".join(lines) + "\n"


def write_report(results_dir: Path):
    results_dir = Path(results_dir)
    meta = json.loads((results_dir / "meta.json").read_text(encoding="utf-8"))
    rows = [json.loads(line) for line in (results_dir / "raw_results.jsonl").read_text(encoding="utf-8").splitlines() if line.strip()]
    text = render_report(meta, rows, load_ratings(results_dir))
    (results_dir / "report.md").write_text(text, encoding="utf-8")
    return text
